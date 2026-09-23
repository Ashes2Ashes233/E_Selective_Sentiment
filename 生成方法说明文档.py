
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "E题三模态选择性情感预测方法与服务器迁移说明.docx"
MD = ROOT / "E题三模态选择性情感预测方法与服务器迁移说明.md"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = tcPr.find(qn("w:shd"))
    if el is None:
        el = OxmlElement("w:shd"); tcPr.append(el)
    el.set(qn("w:fill"), fill)

def set_cell(cell, value, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(value)); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def make_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, "FFFFFF"); shade(t.rows[0].cells[i], "1F4E78")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if ri % 2: shade(cells[i], "F4F7FB")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.25)
    r = p.add_run(text); r.font.name = "Consolas"; r.font.size = Pt(9)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
for name, size, color in [("Normal",10.5,"000000"),("Title",22,"1F4E78"),("Heading 1",15,"1F4E78"),("Heading 2",12,"2F75B5")]:
    st = doc.styles[name]
    st.font.name = "等线"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("E题三模态选择性情感预测方法与服务器迁移说明")
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("基于 Visual-Selective-VIO 思想的比赛实现说明").italic = True
para(doc, "本文说明当前工作区已经实现并运行过的 E 题方法，便于队伍成员理解模型、撰写论文以及把代码迁移到 Linux 或 GPU 服务器上继续训练。维度和路径以当前代码为准；若修改模型配置，需要同步更新检查点和结果。")
make_table(doc, ["项目","当前值"], [
    ["代码目录","E题/selective_sentiment/"],
    ["数据版本","附件2 aligned_50.pkl；对齐版，序列上限50"],
    ["训练/验证/测试样本数","3395 / 728 / 727"],
    ["默认门控","soft（三模态独立软门控）"],
    ["当前最佳检查点","outputs/best.pt，训练第4轮"],
    ["已实测测试集","Accuracy 0.6768；Macro-F1 0.6209；MAE 0.6478；Pearson 0.6637"],
])

doc.add_heading("一 方法总览", level=1)
para(doc, "方法继承母本论文“当前模态信息与历史状态共同决定是否使用某个模态，再进行时序融合”的思路，但把原来的视觉/IMU位姿回归改造成文本、语音、视觉三模态情感预测。整体流程分为问题1的特征与时间组织、问题2的缺失感知预测、问题3的干预式解释三部分。")
code(doc, "输入样本 → 文本/语音/视觉特征 → 可用性掩码 → 三模态门控 → 循环时序融合 → 序列聚合 → 分类头 + 回归头")

doc.add_heading("二 架构结构与输入输出", level=1)
doc.add_heading("2.1 问题1 原始视频特征与对齐", level=2)
para(doc, "问题1负责把100条原始视频整理成可核验的多模态时序特征，不把这100条样本再混入附件2训练集。文本使用题目给定转写；音频统一解码为16 kHz单声道；视频读取帧率、时长和人脸区域。")
make_table(doc, ["阶段","输入","处理","输出"], [
    ["文本","英文转写文本","BERT WordPiece，最多50个位置","token id、attention mask、字符偏移；BERT最后层50×768"],
    ["语音","16 kHz波形","Wav2Vec2 CTC逐帧发射与转写强制对齐；按词元区间均值聚合","50×32声学特征"],
    ["视觉","MP4视频帧","面部区域或整帧的灰度、DCT、HSV、位置、光流统计；按同一时间区间聚合","50×35视觉特征"],
    ["对齐","单词时间区间与WordPiece字符偏移","将语音、视觉帧映射到WordPiece位置；特殊词元和padding置零","valid掩码50；每位置起止秒数"],
])
para(doc, "问题1自生成的声学32维和视觉35维是快速、可解释的代理特征，不能声称与附件2原始74维/35维特征完全同分布。特征、有效掩码、时间区间和处理状态保存在 outputs/q1。")

doc.add_heading("2.2 问题2和问题3预测主干", level=2)
para(doc, "训练和专项预测统一使用附件2对齐版接口。附件2与附件3部分文件缺少连续文本特征，因此代码统一读取 text_bert 的词元编号、注意力掩码和分段编号，再通过同一份冻结BERT生成文本序列表示。音频和视觉直接读取题目提供的特征。")
make_table(doc, ["模块","输入形状（单样本）","输出形状（单样本）","含义"], [
    ["文本编码器","3×50（token id、attention mask、type id）","50×768","每个WordPiece位置的BERT表示；BERT不参与情感任务反向更新"],
    ["语音输入","50×74","50×74","附件2声学序列；按训练集有效位置统计标准化"],
    ["视觉输入","50×35","50×35","附件2视觉序列；按训练集有效位置统计标准化"],
    ["可用性接口","valid：50；availability：50×3","50×3","padding/真实位置以及文本、语音、视觉是否可观测"],
    ["三路适配器","50×768、50×74、50×35","三路50×96","将不同模态转换到同一门控接口"],
    ["门控策略","当前位置三模态表示、三路可用性、历史状态","50×3","文本、语音、视觉各一个保留强度；不可用模态强制为0"],
    ["循环融合","每位置门控特征与可用性","50个时序状态","处理跨位置上下文；padding位置保持上一状态"],
    ["序列聚合","50个时序状态与valid","一个片段级表示","只在有效位置上加权聚合"],
    ["分类头","片段级表示","3个logits和3个概率","Negative、Neutral、Positive"],
    ["回归头","片段级表示","1个标量","情感强度，范围[-3,3]"],
])
para(doc, "批量输入在最前面增加批量维B，例如文本B×50×768、语音B×50×74、视觉B×50×35；预测输出为B×3分类logits、B×3分类概率和B个强度值。")

doc.add_heading("2.3 缺失训练与专项推理", level=2)
para(doc, "训练时在有效序列内随机选择一个或多个模态，并遮挡连续区间。文本在BERT编码之前遮挡，语音和视觉特征直接置零，同时更新availability。模型学习到某一模态不可用时依靠剩余模态和历史状态完成预测。")
para(doc, "附件3的30个对齐版文件用于问题2专项预测；附件4的20个对齐版文件用于问题3专项预测。专项集没有标签，只生成预测和解释文件，不计算专项准确率。")

doc.add_heading("2.4 问题3解释架构", level=2)
para(doc, "解释不是直接把门控值当成因果贡献，而是采用重新推理的遮挡干预。先得到完整输入预测，再分别遮挡文本、语音、视觉，比较目标类别概率和回归强度变化；然后对每个模态逐个遮挡连续4个位置的窗口，选择影响最大的窗口。")
code(doc, "Δ(m,c) = p_c(完整输入) - p_c(遮挡模态m)\nΔ(m,y) = yhat(完整输入) - yhat(遮挡模态m)")
para(doc, "正值表示信息被移除后目标预测下降，通常可解释为支持；负值表示移除后目标预测上升，表示抑制。门控图、模态干预、局部窗口干预、词元区间和视频关键帧分别导出到 outputs/cards、outputs/q3_window_effects.csv、outputs/q3_alignment 和 outputs/keyframes。")

doc.add_heading("三 网络输入输出维度汇总", level=1)
para(doc, "下面只列数据接口和模块接口，不展开策略网络内部的中间层细节。当前检查点使用公共模态维度96和循环状态配置192；若服务器上修改这些配置，必须重新训练并重新生成best.pt。")
make_table(doc, ["对象","输入","输出","当前配置/备注"], [
    ["单条样本输入","text_bert 3×50；audio 50×74；vision 50×35","三模态可用性与valid掩码","对齐版，50为最大位置数"],
    ["BERT文本编码","3×50整数张量","50×768浮点张量","冻结；不计入任务头参数"],
    ["模态适配器","每位置768/74/35","每位置96","三路独立映射"],
    ["门控模块","三路96维表示、可用性、历史状态","每位置3个门控值","文本/语音/视觉各一个；不可用位置为0"],
    ["情感预测器","50位置序列及valid","3类概率和[-3,3]强度","片段级输出"],
    ["解释器","完整输入与遮挡输入","模态差值、窗口差值、时间区间和关键帧","重新执行编码、门控和预测"],
    ["训练批次","B×50×768、B×50×74、B×50×35","B×3 logits、B×3概率、B个强度","当前默认batch size 32"],
])

doc.add_heading("四 评价指标含义", level=1)
para(doc, "题目要求分类任务采用Accuracy、F1，回归任务采用MAE、Pearson。代码另外输出每类F1、加权F1和混淆矩阵，便于分析中性类别。")
make_table(doc, ["指标","公式/定义","含义","使用注意"], [
    ["Accuracy","正确分类数 / 样本总数","整体分类正确比例","类别不平衡时可能被多数类抬高"],
    ["Precision_c","TP_c / (TP_c + FP_c)","预测为c的样本中有多少是真的c","关注误报"],
    ["Recall_c","TP_c / (TP_c + FN_c)","真实c被识别出的比例","关注漏报"],
    ["F1_c","2·Precision_c·Recall_c / (Precision_c + Recall_c)","Precision与Recall的调和平均","同时考虑误报和漏报"],
    ["Macro-F1","三个类别F1的算术平均","每类同等重要的总体分类性能","当前主分类指标；Neutral单独报告"],
    ["Weighted-F1","按各类真实样本数加权的F1平均","兼顾类别比例后的总体F1","受Positive多数类影响较大"],
    ["MAE","mean(|y - yhat|)","预测强度与真实强度的平均绝对距离","越小越好，量纲与情感强度相同"],
    ["Pearson r","cov(y,yhat)/(std(y)·std(yhat))","预测强度与真实强度的线性同向程度","范围[-1,1]，相关性高不等于误差小"],
    ["混淆矩阵","真实类别×预测类别计数表","显示各类互相误判方向","行是真实，列是预测"],
    ["缺失退化量","指标(完整) - 指标(缺失)","衡量局部缺失造成的性能损失","MAE差值、F1差值越接近0越稳健"],
    ["门控使用率","有效位置上的平均门控值","模型内部选择行为的描述","不是官方性能指标，也不是因果贡献百分比"],
])
para(doc, "当前实测：验证集 Accuracy=0.6250、Macro-F1=0.5918、MAE=0.5963、Pearson=0.6543；测试集 Accuracy=0.6768、Macro-F1=0.6209、MAE=0.6478、Pearson=0.6637。测试集Neutral类F1为0.3986，低于Negative和Positive，论文中应单独说明中性边界较难。")
para(doc, "训练选择准则 J=(MAE_clean+MAE_missing)/2 + 0.15×(2-F1_clean-F1_missing)。它不是题目最终指标，而是为了同时考虑完整和缺失输入的模型选择分数。")

doc.add_heading("五 是否可以迁移到服务器上长训练", level=1)
para(doc, "可以。当前代码没有依赖Windows专用API，核心训练脚本可以迁移到Linux GPU服务器。服务器不需要迁移KITTI数据或母本VIO权重，只需要当前selective_sentiment代码、E题数据、BERT权重和训练环境。")
doc.add_heading("5.1 必须携带的内容", level=2)
make_table(doc, ["类别","必须携带","用途"], [
    ["代码","selective_sentiment/*.py；requirements-runtime.txt","训练、推理、问题1特征、问题3解释、打包"],
    ["训练数据","附件2-数据集特征文件/aligned_50.pkl","训练、验证、测试；约1GB量级"],
    ["文本权重","pretrained/bert/","text_bert到50×768；约420MB"],
    ["任务检查点","outputs/best.pt（可选）","继续训练或直接专项推理；约1.7MB"],
    ["训练记录","training_config.json、training_history.csv、input_audit.json","复现实验参数和输入审计"],
    ["问题2专项","附件3对齐版PKL","生成30条缺失预测"],
    ["问题3专项","附件4对齐版PKL及videos/","生成20条解释、时间和关键帧"],
    ["问题1/定位","附件1 MP4、label-100.xlsx、pretrained/aligner/","原始特征、CTC对齐、视频关键帧；ASR约360MB"],
])
para(doc, "如果服务器只做问题2/3长训练，可不带附件1视频、label-100.xlsx、pretrained/aligner和问题1特征生成部分；如果要重新生成问题1或问题3视频证据，则必须带上。")
doc.add_heading("5.2 推荐目录", level=2)
code(doc, "/workspace/E题/selective_sentiment/\n├── common.py model.py train.py predict.py media.py make_report.py\n├── requirements-runtime.txt\n├── pretrained/bert/  pretrained/aligner/\n├── outputs/best.pt\n└── data/E题数据/E题数据/\n    ├── 附件2-数据集特征文件/aligned_50.pkl\n    ├── 附件3-模态缺失特征样本/\n    ├── 附件4-可解释专项视频样本与特征文件/\n    └── 附件1-原始多模态样本/")
doc.add_heading("5.3 Linux安装与运行", level=2)
code(doc, "# 创建环境\npython3.12 -m venv .venv\nsource .venv/bin/activate\n\n# 先按服务器驱动选择兼容CUDA版PyTorch，再安装其余依赖\npip install torch==2.7.1 --index-url https://download.pytorch.org/whl/cu128\npip install -r requirements-runtime.txt --no-deps\n\n# 无网络服务器：提前复制两个pretrained目录\nexport HF_HUB_OFFLINE=1\n\n# 长训练\npython train.py --data /workspace/E题数据/E题数据 \\\n  --output outputs_long --epochs 50 --batch-size 64 --gate soft\n\n# 专项预测和报告\npython predict.py --data /workspace/E题数据/E题数据 \\\n  --checkpoint outputs_long/best.pt --output outputs_long\npython make_report.py --output outputs_long")
para(doc, "迁移时应使用--data、--output、--bert显式指定路径，避免Windows盘符残留。训练脚本会抽查前40个训练样本的BERT tokenizer是否与题目token id匹配；不匹配时主动停止。")
doc.add_heading("5.4 长训练建议", level=2)
for x in [
    "先运行 smoke_check.py，确认文本遮挡、缺失门控、padding不变性和全缺失输入均通过。",
    "保存多个随机种子的 outputs_long_seedXXXX；最后只用验证集综合准则选择模型，不能根据专项测试结果调参。",
    "8GB显存从batch size 32开始；24GB显存可尝试64或96，同时观察显存占用。",
    "允许时将epochs从8提高到30至80，patience提高到8至12；保存训练曲线。",
    "服务器驱动必须兼容CUDA和PyTorch。requirements-runtime.txt中的cu128是本次本地版本，不应机械安装到不兼容服务器。",
    "至少保留训练配置、最佳epoch、验证指标、测试一次结果和最终checkpoint；不必提交所有中间checkpoint。"
]: bullet(doc, x)

doc.add_heading("六 结果文件与论文关系", level=1)
make_table(doc, ["论文内容","对应文件"], [
    ["问题1特征汇总","outputs/q1/sample_summary.csv"],
    ["问题1特征和时间区间","outputs/q1/features/*.npz；outputs/q1/alignment/*.json"],
    ["问题1典型对齐","outputs/q1/alignment_example.csv；example_frame.jpg"],
    ["问题2专项预测","outputs/q2_predictions.csv"],
    ["问题2指标和缺失敏感性","outputs/metrics.json；robustness.csv；figures/robustness.png"],
    ["问题3预测和模态贡献","outputs/q3_predictions_explanations.csv"],
    ["问题3局部证据","outputs/q3_window_effects.csv；cards/；keyframes/"],
    ["复现说明","README.md；模型说明.md；结果说明.md；requirements-runtime.txt"],
])

doc.add_heading("七 论文中应主动说明的边界", level=1)
for x in [
    "当前模型来自单次短训练，不表示全局最优；长训练和多随机种子可能改变结果。",
    "问题1自生成的32/35维特征与附件2的74/35维特征不是同一特征抽取器，不能混写。",
    "门控和注意力是内部选择行为；解释使用遮挡干预，不把门控直接写成因果贡献。",
    "原附件特征没有官方逐词时间戳，问题3定位使用CTC和WordPiece映射，应称为可回看的近似证据定位。",
    "模型依赖文本较多；文本中段缺失70%时性能下降明显，因此只能说具有缺失处理机制，不能夸大为三模态完全均衡鲁棒。"
]: bullet(doc, x)

doc.add_heading("八 来源", level=1)
para(doc, "1. Yang, M., Chen, Y., Kim, H.-S. Efficient Deep Visual and Inertial Odometry with Adaptive Visual Modality Selection. ECCV 2022. https://arxiv.org/abs/2205.06187")
para(doc, "2. 原作者代码：https://github.com/mingyuyng/Visual-Selective-VIO")
para(doc, "3. 工作区E题题面、附件2至附件4数据，以及当前selective_sentiment代码和输出文件。")
para(doc, "4. BERT模型：https://huggingface.co/google-bert/bert-base-uncased；语音CTC模型：https://huggingface.co/facebook/wav2vec2-base-960h")
doc.save(OUT)

MD.write_text("""# E题三模态选择性情感预测方法与服务器迁移说明

完整可编辑版本见同目录DOCX。模型使用附件2对齐版，训练/验证/测试为3395/728/727。输入为text_bert(3×50)、audio(50×74)、vision(50×35)，BERT输出50×768；三路适配到96维后动态门控，循环融合和有效位置聚合，输出三分类概率与[-3,3]情感强度。

## 当前实测

测试集 Accuracy=0.6768，Macro-F1=0.6209，MAE=0.6478，Pearson=0.6637。验证集 Accuracy=0.6250，Macro-F1=0.5918，MAE=0.5963，Pearson=0.6543。

## 服务器迁移

可以迁移。必须带selective_sentiment/*.py、requirements-runtime.txt、附件2 aligned_50.pkl、pretrained/bert/和可选的outputs/best.pt。若还要跑问题1和问题3视频定位，再带附件1视频、label-100.xlsx、附件3/4、pretrained/aligner/。Linux上先按服务器CUDA驱动安装兼容PyTorch，再安装其余依赖；设置HF_HUB_OFFLINE=1时需提前复制两套预训练权重。训练示例：python train.py --data /workspace/E题数据/E题数据 --output outputs_long --epochs 50 --batch-size 64 --gate soft。

完整架构、维度、指标公式、目录和论文边界见DOCX。
""", encoding="utf-8")
print(OUT)

